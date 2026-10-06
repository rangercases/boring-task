import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def analyze_assortment_file(file_path, output_path=None):
    """
    Phân tích file Assortment theo từng tuần PO và điểm đến (CW01, CW02).
    Tạo hoặc cập nhật sheet 'Per_week' với đầy đủ công thức và định dạng.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

    # Mở workbook gốc để giữ nguyên các sheet khác
    wb = openpyxl.load_workbook(file_path)
    
    # Tìm sheet dữ liệu nguồn (ưu tiên 'Details', hoặc sheet đầu tiên)
    source_sheet_name = 'Details' if 'Details' in wb.sheetnames else wb.sheetnames[0]
    ws_source = wb[source_sheet_name]

    # Đọc headers dòng 1
    headers = {}
    for col in range(1, ws_source.max_column + 1):
        val = ws_source.cell(1, col).value
        if val is not None:
            clean_name = str(val).strip().lower()
            headers[clean_name] = col

    # Nhận diện các cột cần thiết (hỗ trợ nhiều biến thể tên cột)
    po_col = None
    for name in ['po.', 'po', 'po number', 'tuần', 'week']:
        if name in headers:
            po_col = headers[name]
            break

    dest_col = None
    for name in ['destination', 'dest', 'kho', 'kho đích']:
        if name in headers:
            dest_col = headers[name]
            break

    sku_col = None
    for name in ['no.', 'no', 'item no', 'sku', 'mã hàng', 'item code']:
        if name in headers:
            sku_col = headers[name]
            break

    if not po_col or not dest_col or not sku_col:
        raise ValueError("Không tìm thấy đủ các cột bắt buộc: PO (hoặc PO.), Destination, và No. (Mã hàng) trong sheet dữ liệu.")

    # Đọc dữ liệu và gom nhóm theo PO
    # Cấu trúc: po_groups[po_name] = {'CW01': set(sku), 'CW02': set(sku)}
    po_groups = {}
    # Thứ tự xuất hiện ban đầu của PO để giữ đúng layout
    po_order = []

    for r in range(2, ws_source.max_row + 1):
        po_val = ws_source.cell(r, po_col).value
        dest_val = ws_source.cell(r, dest_col).value
        sku_val = ws_source.cell(r, sku_col).value

        if po_val is None or sku_val is None:
            continue

        po_str = str(po_val).strip()
        if not po_str:
            continue

        dest_str = str(dest_val).strip().upper() if dest_val is not None else ""
        sku_str = str(sku_val).strip()

        if po_str not in po_groups:
            po_groups[po_str] = {"CW01": set(), "CW02": set()}
            po_order.append(po_str)

        if "CW01" in dest_str:
            po_groups[po_str]["CW01"].add(sku_str)
        elif "CW02" in dest_str:
            po_groups[po_str]["CW02"].add(sku_str)

    if not po_order:
        raise ValueError("Không tìm thấy dòng dữ liệu hợp lệ nào để phân tích.")

    # Xử lý sheet 'Per_week'
    if 'Per_week' in wb.sheetnames:
        del wb['Per_week']
    ws_out = wb.create_sheet(title='Per_week')

    # Định nghĩa Styles Apple UI
    header_fill = PatternFill(start_color="3B302A", end_color="3B302A", fill_type="solid") # Dark Roast Brown
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    bold_data_font = Font(name="Segoe UI", size=10, bold=True)
    pct_font = Font(name="Segoe UI", size=10, bold=True, color="2E5A36") # Xanh lá đậm cho tỷ lệ

    thin_border_side = Side(border_style="thin", color="E0D6C8")
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Nền xen kẽ dịu mắt
    zebra_fill = PatternFill(start_color="FAF7F2", end_color="FAF7F2", fill_type="solid")
    highlight_fill = PatternFill(start_color="F2EBD9", end_color="F2EBD9", fill_type="solid")

    columns = [
        "PO",
        "CW01",
        "CW02",
        "Duplicated",
        "Only CW01",
        "Only CW02",
        "SKU/week",
        "% Denmark Codes Also in Belgium"
    ]

    # Ghi Header
    for c_idx, col_name in enumerate(columns, 1):
        cell = ws_out.cell(row=1, column=c_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws_out.row_dimensions[1].height = 28

    # Ghi dữ liệu từng tuần
    for r_idx, po_name in enumerate(po_order, 2):
        s_cw01 = po_groups[po_name]["CW01"]
        s_cw02 = po_groups[po_name]["CW02"]

        cnt_cw01 = len(s_cw01)
        cnt_cw02 = len(s_cw02)
        duplicated = len(s_cw01 & s_cw02)
        only_cw01 = len(s_cw01 - s_cw02)
        only_cw02 = len(s_cw02 - s_cw01)

        # Ghi các giá trị
        c_po = ws_out.cell(row=r_idx, column=1, value=po_name)
        c_cw01 = ws_out.cell(row=r_idx, column=2, value=cnt_cw01)
        c_cw02 = ws_out.cell(row=r_idx, column=3, value=cnt_cw02)
        c_dup = ws_out.cell(row=r_idx, column=4, value=duplicated)
        c_o1 = ws_out.cell(row=r_idx, column=5, value=only_cw01)
        c_o2 = ws_out.cell(row=r_idx, column=6, value=only_cw02)
        
        # Công thức tính Tổng SKU và Tỷ lệ theo đúng form của Nhung
        c_sku = ws_out.cell(row=r_idx, column=7, value=f"=SUM(D{r_idx}:F{r_idx})")
        c_pct = ws_out.cell(row=r_idx, column=8, value=f"=IF(B{r_idx}>0, D{r_idx}/B{r_idx}, 0)")
        c_pct.number_format = '0.0%'

        is_even = (r_idx % 2 == 0)
        curr_fill = zebra_fill if is_even else PatternFill(fill_type=None)

        # Căn chỉnh và border
        for c_idx in range(1, 9):
            cell = ws_out.cell(row=r_idx, column=c_idx)
            cell.border = data_border
            if c_idx == 1:
                cell.font = data_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx == 4: # Duplicated highlight
                cell.font = bold_data_font
                cell.fill = highlight_fill
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 7: # SKU/week
                cell.font = bold_data_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 8: # %
                cell.font = pct_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.font = data_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            
            if c_idx not in (4,) and curr_fill.fill_type:
                cell.fill = curr_fill

        ws_out.row_dimensions[r_idx].height = 22

    # Tự động điều chỉnh độ rộng cột
    for col in ws_out.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_out.column_dimensions[col_letter].width = max(max_len + 4, 12)
    
    # Cột PO và Cột % cho rộng rãi hơn
    ws_out.column_dimensions['A'].width = 30
    ws_out.column_dimensions['H'].width = 32

    # Lưu file kết quả
    if not output_path:
        base, ext = os.path.splitext(file_path)
        output_path = f"{base}_analyzed{ext}"

    wb.save(output_path)
    wb.close()
    return output_path, len(po_order)
