import os
import sys
import datetime
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter

class ClaimHelperApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Claim Helper - SOFACOMPANY Claim & Cost Matcher v1.0")
        self.root.geometry("820x720")
        self.root.minsize(750, 600)
        
        # Color Palette & Styles
        self.bg_color = "#f4f6f9"
        self.primary_color = "#1e3a8a" # Navy Blue
        self.accent_color = "#2563eb"
        self.success_color = "#16a34a"
        self.text_color = "#1f2937"
        
        self.root.configure(bg=self.bg_color)
        
        # Variables
        self.overview_path = tk.StringVar(value="")
        self.target_folder = tk.StringVar(value="")
        self.auto_supplier = tk.BooleanVar(value=True)
        self.add_note_col = tk.BooleanVar(value=True)
        self.add_total_sum = tk.BooleanVar(value=True)
        self.filter_partials = tk.BooleanVar(value=True)
        
        # Auto-detect default paths if in same directory
        curr_dir = os.path.abspath(os.path.dirname(__file__))
        self.target_folder.set(curr_dir)
        for f in os.listdir(curr_dir):
            if f.lower().startswith("overview") and f.endswith(".xlsx"):
                self.overview_path.set(os.path.join(curr_dir, f))
                break

        self.setup_ui()

    def setup_ui(self):
        # Header Banner
        header_frame = tk.Frame(self.root, bg=self.primary_color, height=80)
        header_frame.pack(fill="x", side="top")
        
        title_label = tk.Label(
            header_frame, 
            text="CLAIM HELPER", 
            font=("Segoe UI", 18, "bold"), 
            fg="white", 
            bg=self.primary_color
        )
        title_label.pack(anchor="w", padx=20, pady=(12, 2))
        
        subtitle_label = tk.Label(
            header_frame, 
            text="Hệ thống tự động đối chiếu & điền Purchase Cost cho các file khiếu nại (Claim Versus)", 
            font=("Segoe UI", 9), 
            fg="#93c5fd", 
            bg=self.primary_color
        )
        subtitle_label.pack(anchor="w", padx=20, pady=(0, 12))

        # Main Scrollable / Content Container
        main_container = tk.Frame(self.root, bg=self.bg_color)
        main_container.pack(fill="both", expand=True, padx=20, pady=15)

        # Section 1: Master Overview File
        sec1 = tk.LabelFrame(
            main_container, 
            text=" 1. File Master Overview ", 
            font=("Segoe UI", 10, "bold"), 
            fg=self.primary_color, 
            bg=self.bg_color,
            padx=12, pady=10
        )
        sec1.pack(fill="x", pady=(0, 10))

        e1 = tk.Entry(sec1, textvariable=self.overview_path, font=("Segoe UI", 9), bg="white", relief="solid", bd=1)
        e1.pack(side="left", fill="x", expand=True, padx=(0, 10), ipady=4)

        b1 = tk.Button(
            sec1, text="Chọn File Overview...", 
            font=("Segoe UI", 9, "bold"), 
            bg=self.accent_color, fg="white", 
            relief="flat", cursor="hand2", padx=12, pady=3,
            command=self.browse_overview
        )
        b1.pack(side="right")

        # Section 2: Target Versus Folder
        sec2 = tk.LabelFrame(
            main_container, 
            text=" 2. Thư mục chứa các file Claim Versus ", 
            font=("Segoe UI", 10, "bold"), 
            fg=self.primary_color, 
            bg=self.bg_color,
            padx=12, pady=10
        )
        sec2.pack(fill="x", pady=(0, 10))

        e2 = tk.Entry(sec2, textvariable=self.target_folder, font=("Segoe UI", 9), bg="white", relief="solid", bd=1)
        e2.pack(side="left", fill="x", expand=True, padx=(0, 10), ipady=4)

        b2 = tk.Button(
            sec2, text="Chọn Thư Mục...", 
            font=("Segoe UI", 9, "bold"), 
            bg=self.accent_color, fg="white", 
            relief="flat", cursor="hand2", padx=12, pady=3,
            command=self.browse_folder
        )
        b2.pack(side="right")

        # Section 3: Options Frame
        sec3 = tk.LabelFrame(
            main_container, 
            text=" 3. Cấu hình đối chiếu thông minh ", 
            font=("Segoe UI", 10, "bold"), 
            fg=self.primary_color, 
            bg=self.bg_color,
            padx=12, pady=8
        )
        sec3.pack(fill="x", pady=(0, 10))

        opt_grid = tk.Frame(sec3, bg=self.bg_color)
        opt_grid.pack(fill="x")

        tk.Checkbutton(opt_grid, text="Tự động nhận diện Nhà cung cấp (Casa / Nhan Hoang)", variable=self.auto_supplier, bg=self.bg_color, font=("Segoe UI", 9)).grid(row=0, column=0, sticky="w", padx=10, pady=2)
        tk.Checkbutton(opt_grid, text="Lọc bỏ phụ kiện/kiện lẻ (như 'Only box 1 of 2')", variable=self.filter_partials, bg=self.bg_color, font=("Segoe UI", 9)).grid(row=0, column=1, sticky="w", padx=10, pady=2)
        tk.Checkbutton(opt_grid, text="Thêm cột ghi chú giải thích (Purchase Cost Note)", variable=self.add_note_col, bg=self.bg_color, font=("Segoe UI", 9)).grid(row=1, column=0, sticky="w", padx=10, pady=2)
        tk.Checkbutton(opt_grid, text="Thêm công thức Tổng cộng động (=SUM)", variable=self.add_total_sum, bg=self.bg_color, font=("Segoe UI", 9)).grid(row=1, column=1, sticky="w", padx=10, pady=2)

        # Section 4: Action & Progress
        act_frame = tk.Frame(main_container, bg=self.bg_color)
        act_frame.pack(fill="x", pady=(5, 10))

        self.btn_run = tk.Button(
            act_frame, 
            text="▶ BẮT ĐẦU ĐỐI CHIẾU & ĐIỀN GIÁ", 
            font=("Segoe UI", 11, "bold"), 
            bg=self.success_color, fg="white", 
            relief="flat", cursor="hand2", padx=20, pady=8,
            command=self.start_processing
        )
        self.btn_run.pack(side="left", padx=(0, 15))

        self.btn_open_folder = tk.Button(
            act_frame, 
            text="📂 Mở Thư Mục Chứa File", 
            font=("Segoe UI", 10), 
            bg="#6b7280", fg="white", 
            relief="flat", cursor="hand2", padx=15, pady=8,
            state="normal",
            command=self.open_target_folder
        )
        self.btn_open_folder.pack(side="left")

        self.progress_bar = ttk.Progressbar(main_container, mode="indeterminate")
        self.progress_bar.pack(fill="x", pady=(0, 8))

        # Section 5: Real-time Log Console
        sec5 = tk.LabelFrame(
            main_container, 
            text=" Nhật ký xử lý chi tiết ", 
            font=("Segoe UI", 9, "bold"), 
            fg="#4b5563", 
            bg=self.bg_color,
            padx=10, pady=8
        )
        sec5.pack(fill="both", expand=True)

        self.log_area = scrolledtext.ScrolledText(
            sec5, 
            wrap="word", 
            font=("Consolas", 9), 
            bg="#111827", fg="#f3f4f6", 
            relief="flat", padx=8, pady=8
        )
        self.log_area.pack(fill="both", expand=True)
        
        self.log("Chào mừng bạn đến với Claim Helper!")
        self.log("Vui lòng kiểm tra file Overview và thư mục Claim Versus, sau đó bấm 'BẮT ĐẦU ĐỐI CHIẾU'.\n")

    def log(self, message):
        self.log_area.insert("end", f"[{datetime.datetime.now().strftime('%H:%M:%S')}] {message}\n")
        self.log_area.see("end")

    def browse_overview(self):
        f = filedialog.askopenfilename(
            title="Chọn file Master Overview",
            filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")]
        )
        if f:
            self.overview_path.set(f)

    def browse_folder(self):
        d = filedialog.askdirectory(title="Chọn thư mục chứa các file Claim Versus")
        if d:
            self.target_folder.set(d)

    def open_target_folder(self):
        d = self.target_folder.get().strip()
        if os.path.exists(d):
            os.startfile(d)
        else:
            messagebox.showwarning("Cảnh báo", "Đường dẫn thư mục không tồn tại!")

    def start_processing(self):
        ov_path = self.overview_path.get().strip()
        tg_dir = self.target_folder.get().strip()

        if not os.path.isfile(ov_path):
            messagebox.showerror("Lỗi", "Vui lòng chọn file Master Overview hợp lệ (.xlsx)!")
            return
        if not os.path.isdir(tg_dir):
            messagebox.showerror("Lỗi", "Vui lòng chọn thư mục chứa các file Claim Versus!")
            return

        # Disable button & start progress bar
        self.btn_run.config(state="disabled", text="⏳ ĐANG XỬ LÝ DỮ LIỆU...")
        self.progress_bar.start(10)
        self.log("="*65)
        self.log("BẮT ĐẦU QUY TRÌNH ĐỐI CHIẾU VÀ ĐIỀN GIÁ")
        self.log(f"Overview: {os.path.basename(ov_path)}")
        self.log(f"Thư mục Versus: {tg_dir}")

        # Run background thread
        thread = threading.Thread(target=self.run_matcher_core, args=(ov_path, tg_dir), daemon=True)
        thread.start()

    def run_matcher_core(self, overview_path, target_dir):
        try:
            # 1. Load Overview
            self.log("\n[1/3] Đang nạp dữ liệu từ file Master Overview (vui lòng chờ vài giây)...")
            wb_o = openpyxl.load_workbook(overview_path, read_only=True, data_only=True)
            s_o = wb_o['Overview'] if 'Overview' in wb_o.sheetnames else wb_o.active
            headers_o = next(s_o.iter_rows(max_row=1, values_only=True))
            col_idx = {h: i for i, h in enumerate(headers_o)}

            cs_records = []
            nh_records = []

            for r in s_o.iter_rows(min_row=2, values_only=True):
                man = str(r[col_idx.get('Manufacturer', 0)] or '').strip().upper()
                sup = str(r[col_idx.get('Supplier', 0)] or '').strip().upper()
                cost = r[col_idx.get('Purchase cost', 0)]
                note_p = str(r[col_idx.get('Item PO note', 0)] or '').lower()
                note_s = str(r[col_idx.get('Note (sales)', 0)] or '').lower()

                if self.filter_partials.get():
                    if 'box 1 of 2' in note_p or 'box 2 of 2' in note_p or 'box 1 of' in note_s:
                        continue

                if cost is None or float(cost) <= 0:
                    continue

                rec = {
                    'art': str(r[col_idx.get('Customer art No', 0)] or '').strip(),
                    'item_name': str(r[col_idx.get('Customer item name', 0)] or '').strip(),
                    'fabric': str(r[col_idx.get('Fabric type', 0)] or '').strip(),
                    'cost': float(cost),
                    'po_date': r[col_idx.get('P.O issued date', 0)],
                    'model': str(r[col_idx.get('Model', 0)] or '').strip()
                }

                if man in ['CS', 'WORKSHOP'] or 'CASA' in man or 'CASA' in sup:
                    cs_records.append(rec)
                elif man in ['NH', 'NHF', 'NH FSC', 'NHAN HOANG'] or 'NHAN HOANG' in sup or 'NH ' in sup:
                    nh_records.append(rec)

            wb_o.close()
            self.log(f"-> Nạp xong Master Data: {len(cs_records)} bản ghi Casa (CS) | {len(nh_records)} bản ghi Nhan Hoang (NH).")

            # 2. Find Versus Files
            self.log("\n[2/3] Quét các file Claim Versus trong thư mục...")
            target_files = [
                f for f in os.listdir(target_dir) 
                if f.endswith('.xlsx') 
                and 'versus' in f.lower() 
                and not f.endswith('_backup.xlsx')
                and not f.startswith('~$')
            ]

            if not target_files:
                self.log("❌ Không tìm thấy file Excel nào có từ khóa 'versus' trong thư mục!")
                self.root.after(0, self.finish_processing, False, "Không tìm thấy file nào để xử lý!")
                return

            self.log(f"-> Tìm thấy {len(target_files)} file cần xử lý:")
            for tf in target_files:
                self.log(f"   • {tf}")

            # 3. Process Each File
            self.log("\n[3/3] Bắt đầu xử lý từng file...")
            total_claims_all = 0
            total_matched_all = 0

            for tf in target_files:
                filepath = os.path.join(target_dir, tf)
                is_casa = 'casa' in tf.lower()
                ref_db = cs_records if is_casa else nh_records
                supplier_name = "Casa (CS)" if is_casa else "Nhan Hoang (NH)"

                self.log(f"\n--- Đang xử lý: {tf} [{supplier_name}] ---")
                wb = openpyxl.load_workbook(filepath)
                s = wb['Export'] if 'Export' in wb.sheetnames else wb.active
                headers = [cell.value for cell in s[1]]

                # Add column headers
                if 'Purchase Cost (USD)' in headers:
                    col_cost_idx = headers.index('Purchase Cost (USD)') + 1
                else:
                    col_cost_idx = len(headers) + 1
                    headers.append('Purchase Cost (USD)')
                    c_h1 = s.cell(row=1, column=col_cost_idx, value='Purchase Cost (USD)')
                    sample_h = s.cell(row=1, column=headers.index('Unit Cost x QTY')+1 if 'Unit Cost x QTY' in headers else 1)
                    if sample_h.has_style:
                        c_h1.font, c_h1.border, c_h1.fill, c_h1.alignment = copy(sample_h.font), copy(sample_h.border), copy(sample_h.fill), copy(sample_h.alignment)

                if self.add_note_col.get():
                    if 'Purchase Cost Note' in headers:
                        col_note_idx = headers.index('Purchase Cost Note') + 1
                    else:
                        col_note_idx = len(headers) + 1
                        headers.append('Purchase Cost Note')
                        c_h2 = s.cell(row=1, column=col_note_idx, value='Purchase Cost Note')
                        sample_h = s.cell(row=1, column=headers.index('Unit Cost x QTY')+1 if 'Unit Cost x QTY' in headers else 1)
                        if sample_h.has_style:
                            c_h2.font, c_h2.border, c_h2.fill, c_h2.alignment = copy(sample_h.font), copy(sample_h.border), copy(sample_h.fill), copy(sample_h.alignment)

                cost_col_letter = get_column_letter(col_cost_idx)
                s.column_dimensions[cost_col_letter].width = 22
                if self.add_note_col.get():
                    note_col_letter = get_column_letter(col_note_idx)
                    s.column_dimensions[note_col_letter].width = 85

                max_row = s.max_row
                art_col_idx = headers.index('No') + 1 if 'No' in headers else 7
                unit_cost_col_idx = headers.index('Unit Cost x QTY') + 1 if 'Unit Cost x QTY' in headers else None

                # Check total row
                last_art = s.cell(row=max_row, column=art_col_idx).value
                last_row_is_total = (last_art is None or str(last_art).strip() == '')
                data_end_row = max_row - 1 if last_row_is_total else max_row
                data_rows_count = data_end_row - 1

                # Median rate
                rate_samples = []
                if unit_cost_col_idx:
                    for r_idx in range(2, data_end_row + 1):
                        art = str(s.cell(row=r_idx, column=art_col_idx).value or '').strip()
                        dkk = s.cell(row=r_idx, column=unit_cost_col_idx).value
                        exact = [rec for rec in ref_db if rec['art'] == art]
                        if exact and isinstance(dkk, (int, float)) and dkk > 0:
                            avg_usd = sum(e['cost'] for e in exact) / len(exact)
                            rate_samples.append(dkk / avg_usd)

                median_rate = sorted(rate_samples)[len(rate_samples)//2] if rate_samples else (8.0087 if is_casa else 7.8210)

                file_matched = 0
                for r_idx in range(2, data_end_row + 1):
                    art = str(s.cell(row=r_idx, column=art_col_idx).value or '').strip()
                    desc = str(s.cell(row=r_idx, column=headers.index('Description')+1).value or '').strip() if 'Description' in headers else ''
                    desc2 = str(s.cell(row=r_idx, column=headers.index('Description_2')+1).value or '').strip() if 'Description_2' in headers else ''
                    posting_date = s.cell(row=r_idx, column=headers.index('PostingDate')+1).value if 'PostingDate' in headers else None
                    days = s.cell(row=r_idx, column=headers.index('Days Between Order And Complaint')+1).value if 'Days Between Order And Complaint' in headers else 0
                    days = days if isinstance(days, (int, float)) else 0
                    unit_cost_dkk = s.cell(row=r_idx, column=unit_cost_col_idx).value if unit_cost_col_idx else 0
                    unit_cost_dkk = unit_cost_dkk if isinstance(unit_cost_dkk, (int, float)) else 0

                    order_date = None
                    if isinstance(posting_date, datetime.datetime):
                        order_date = posting_date - datetime.timedelta(days=int(days))
                    elif isinstance(posting_date, datetime.date):
                        order_date = posting_date - datetime.timedelta(days=int(days))

                    exact_matches = [rec for rec in ref_db if rec['art'] == art]
                    selected_cost = None
                    note = ""

                    if exact_matches:
                        file_matched += 1
                        if unit_cost_dkk > 0:
                            target_usd = unit_cost_dkk / median_rate
                            plausible = [e for e in exact_matches if 0.7 * target_usd <= e['cost'] <= 1.3 * target_usd]
                            if not plausible: plausible = exact_matches
                        else:
                            plausible = exact_matches

                        valid_dates = [e for e in plausible if isinstance(e['po_date'], (datetime.datetime, datetime.date))]
                        if valid_dates and order_date:
                            o_d = order_date.date() if isinstance(order_date, datetime.datetime) else order_date
                            best_match = min(valid_dates, key=lambda e: abs((e['po_date'].date() if isinstance(e['po_date'], datetime.datetime) else e['po_date']) - o_d))
                            selected_cost = best_match['cost']
                            po_str = best_match['po_date'].strftime('%d/%m/%Y')
                            od_str = order_date.strftime('%d/%m/%Y')
                            note = f"Khớp 100% Art No ({art}). Chọn {selected_cost}$ theo đợt PO ({po_str}) gần nhất ngày đặt hàng ({od_str})."
                        else:
                            best_match = min(plausible, key=lambda e: abs(e['cost'] - (unit_cost_dkk/median_rate if unit_cost_dkk > 0 else plausible[0]['cost'])))
                            selected_cost = best_match['cost']
                            note = f"Khớp 100% Art No ({art}). Chọn {selected_cost}$ từ danh mục trong Overview."
                    else:
                        candidates = [rec for rec in ref_db if rec['art'].startswith(art) or art in rec['art']]
                        if not candidates and len(art) >= 6:
                            candidates = [rec for rec in ref_db if rec['art'].startswith(art[:6])]

                        if candidates and unit_cost_dkk > 0:
                            file_matched += 1
                            est_usd = unit_cost_dkk / median_rate
                            best_cand = min(candidates, key=lambda c: abs(c['cost'] - est_usd))
                            selected_cost = best_cand['cost']
                            note = f"Mã rút gọn MM ({art}). Dựa vào Unit Cost {unit_cost_dkk:.2f} DKK (~{est_usd:.2f}$), khớp phân khúc vải '{best_cand['fabric']}' (giá {selected_cost}$)."
                        elif candidates:
                            file_matched += 1
                            selected_cost = round(sum(c['cost'] for c in candidates) / len(candidates), 2)
                            note = f"Mã rút gọn MM ({art}). Lấy giá trung bình phân khúc {selected_cost}$."
                        elif unit_cost_dkk > 0:
                            file_matched += 1
                            selected_cost = round(unit_cost_dkk / median_rate, 2)
                            note = f"Mã {art} không còn trong Overview. Tính theo tỷ giá quy đổi Unit Cost: {unit_cost_dkk:.2f} DKK / {median_rate:.4f} = {selected_cost}$."
                        else:
                            selected_cost = 0.0
                            note = f"Không tìm thấy mã {art} trong Overview."

                    c_cost = s.cell(row=r_idx, column=col_cost_idx, value=selected_cost)
                    c_cost.number_format = '#,##0.00'
                    sample_c = s.cell(row=r_idx, column=unit_cost_col_idx if unit_cost_col_idx else 1)
                    if sample_c.has_style:
                        c_cost.font, c_cost.border = copy(sample_c.font), copy(sample_c.border)

                    if self.add_note_col.get():
                        c_note = s.cell(row=r_idx, column=col_note_idx, value=note)
                        if sample_c.has_style:
                            c_note.font, c_note.border = copy(sample_c.font), copy(sample_c.border)

                # Total Row
                if last_row_is_total and self.add_total_sum.get():
                    c_total = s.cell(row=max_row, column=col_cost_idx, value=f"=SUM({cost_col_letter}2:{cost_col_letter}{data_end_row})")
                    c_total.number_format = '#,##0.00'
                    sample_t = s.cell(row=max_row, column=unit_cost_col_idx if unit_cost_col_idx else 1)
                    if sample_t.has_style:
                        c_total.font, c_total.border = copy(sample_t.font), copy(sample_t.border)
                    if self.add_note_col.get():
                        c_t_note = s.cell(row=max_row, column=col_note_idx, value=f"Tổng Purchase Cost (USD) của {data_rows_count} vụ khiếu nại")
                        if sample_t.has_style:
                            c_t_note.font, c_t_note.border = copy(sample_t.font), copy(sample_t.border)

                wb.save(filepath)
                wb.close()

                rate_str = f"{(file_matched / data_rows_count * 100):.1f}%" if data_rows_count > 0 else "100%"
                self.log(f"   ✓ Hoàn tất {tf}: {file_matched}/{data_rows_count} dòng khớp ({rate_str})")
                total_claims_all += data_rows_count
                total_matched_all += file_matched

            # Finished all
            self.log("\n" + "="*65)
            self.log(f"TỔNG KẾT: Đã xử lý {len(target_files)} file | Tổng cộng {total_matched_all}/{total_claims_all} lượt claim ({total_matched_all/total_claims_all*100:.1f}%)")
            self.log("="*65)

            self.root.after(0, self.finish_processing, True, f"Xử lý thành công toàn bộ {len(target_files)} file ({total_matched_all} dòng claim)!")

        except Exception as e:
            self.log(f"\n❌ LỖI TRONG QUÁ TRÌNH XỬ LÝ: {str(e)}")
            self.root.after(0, self.finish_processing, False, str(e))

    def finish_processing(self, success, msg):
        self.progress_bar.stop()
        self.btn_run.config(state="normal", text="▶ BẮT ĐẦU ĐỐI CHIẾU & ĐIỀN GIÁ")
        if success:
            messagebox.showinfo("Thành công", msg)
        else:
            messagebox.showerror("Thông báo lỗi", msg)

if __name__ == "__main__":
    root = tk.Tk()
    app = ClaimHelperApp(root)
    root.mainloop()
