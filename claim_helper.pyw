import os
import sys
import re
import datetime
import threading
import pickle
import json
import subprocess
import urllib.request
import urllib.error
import tempfile
import time
from copy import copy

import customtkinter as ctk
from tkinter import filedialog, messagebox
import tkinterdnd2 as tkdnd
import openpyxl
from openpyxl.utils import get_column_letter

# Set Apple-inspired Light Mode
ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")

APP_VERSION = "v1.0.0"
GITHUB_REPO = "rangercases/claim-helper"
CACHE_FILE_NAME = ".overview_cache.pkl"
STATE_FILE_NAME = ".app_state.json"

def parse_version(v_str):
    clean = re.sub(r'[^0-9.]', '', str(v_str))
    parts = []
    for p in clean.split('.'):
        if p.isdigit():
            parts.append(int(p))
    return parts

def is_newer_version(latest_tag, current_tag):
    v_latest = parse_version(latest_tag)
    v_curr = parse_version(current_tag)
    return v_latest > v_curr

def get_cache_path():
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), CACHE_FILE_NAME)

def get_state_path():
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), STATE_FILE_NAME)

def parse_drop_paths(data_str):
    """Parse dropped file paths handling curly braces and spaces."""
    if not data_str:
        return []
    pattern = r'\{([^}]+)\}|(\S+)'
    matches = re.findall(pattern, data_str.strip())
    paths = []
    for m in matches:
        p = (m[0] if m[0] else m[1]).strip().strip('"').strip("'")
        if os.path.exists(p):
            paths.append(os.path.abspath(p))
    return paths

def format_file_size(size_bytes):
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.1f} MB"

def reveal_in_explorer(filepath):
    """Opens Windows Explorer with the specific file selected/highlighted."""
    try:
        norm_path = os.path.normpath(filepath)
        if os.path.exists(norm_path):
            subprocess.Popen(f'explorer /select,"{norm_path}"')
        else:
            folder = os.path.dirname(norm_path)
            if os.path.exists(folder):
                subprocess.Popen(f'explorer "{folder}"')
    except Exception as e:
        print(f"Error revealing in explorer: {e}")

class ClaimHelperAppleApp(ctk.CTk, tkdnd.TkinterDnD.DnDWrapper):
    def __init__(self):
        super().__init__()
        self.TkdndVersion = tkdnd.TkinterDnD._require(self)

        self.title("Claim Helper")
        self.geometry("840x820")
        self.minsize(800, 720)
        self.configure(fg_color="#F5F5F7")  # Signature Apple off-white

        # State Variables
        self.overview_path = ""
        self.claim_files = [] # list of absolute paths
        self.last_created_files = [] # list of created output files
        self.is_processing = False

        curr_dir = os.path.abspath(os.path.dirname(__file__))
        saved_state = self.load_saved_state()

        if saved_state is not None:
            # Restore saved overview path if it still exists
            saved_ov = saved_state.get("overview_path", "")
            if saved_ov and os.path.exists(saved_ov):
                self.overview_path = saved_ov
            
            # Restore saved claim files (keep only existing files)
            saved_claims = saved_state.get("claim_files", [])
            self.claim_files = [p for p in saved_claims if os.path.exists(p)]
        else:
            # First launch: Auto-detect Overview in current directory
            for f in os.listdir(curr_dir):
                if f.lower().startswith("overview") and f.endswith(".xlsx") and not f.startswith("~$"):
                    self.overview_path = os.path.join(curr_dir, f)
                    break

            # First launch: Auto-detect initial versus files in current directory
            for f in os.listdir(curr_dir):
                if "versus" in f.lower() and f.endswith(".xlsx") and not f.endswith("_backup.xlsx") and not f.endswith("_filled.xlsx") and not f.startswith("~$"):
                    self.claim_files.append(os.path.join(curr_dir, f))
            
            self.save_state()

        # If overview not found in saved state, fallback to auto-detecting in current dir
        if not self.overview_path:
            for f in os.listdir(curr_dir):
                if f.lower().startswith("overview") and f.endswith(".xlsx") and not f.startswith("~$"):
                    self.overview_path = os.path.join(curr_dir, f)
                    break

        self.setup_ui()
        self.setup_drag_and_drop()

        # Check for remote updates silently in background
        threading.Thread(target=self.check_for_updates, daemon=True).start()

    def load_saved_state(self):
        """Loads persistent session state if available."""
        s_path = get_state_path()
        if not os.path.exists(s_path):
            return None
        try:
            with open(s_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            return data
        except Exception as e:
            print(f"Error loading state: {e}")
            return None

    def save_state(self):
        """Saves current state to persistent storage."""
        s_path = get_state_path()
        try:
            data = {
                "overview_path": self.overview_path,
                "claim_files": self.claim_files
            }
            with open(s_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"Error saving state: {e}")

    # ----------------------------------------------------
    # GitHub Auto-Updater (Non-blocking & Seamless)
    # ----------------------------------------------------
    def check_for_updates(self):
        """Silently checks GitHub Releases for new versions."""
        try:
            url = f"https://api.github.com/repos/{GITHUB_REPO}/releases/latest"
            req = urllib.request.Request(url, headers={"User-Agent": "ClaimHelper-App"})
            with urllib.request.urlopen(req, timeout=5) as resp:
                if resp.status == 200:
                    data = json.loads(resp.read().decode("utf-8"))
                    latest_tag = data.get("tag_name", "")
                    if latest_tag and is_newer_version(latest_tag, APP_VERSION):
                        assets = data.get("assets", [])
                        exe_url = None
                        for a in assets:
                            if a.get("name", "").lower().endswith(".exe"):
                                exe_url = a.get("browser_download_url")
                                break
                        if exe_url:
                            self.after(0, self.show_update_banner, latest_tag, exe_url)
        except Exception:
            # Offline or GitHub rate limit - keep user experience uninterrupted
            pass

    def show_update_banner(self, latest_tag, exe_url):
        for child in self.update_banner_container.winfo_children():
            child.destroy()

        banner = ctk.CTkFrame(
            self.update_banner_container,
            fg_color="#EBF3FE",
            corner_radius=10,
            border_width=1,
            border_color="#C7DEFF"
        )
        banner.pack(fill="x")

        left = ctk.CTkFrame(banner, fg_color="transparent")
        left.pack(side="left", padx=14, pady=10, fill="x", expand=True)

        ctk.CTkLabel(
            left,
            text=f"🚀 Đã có phiên bản mới ({latest_tag})!",
            font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
            text_color="#0071E3"
        ).pack(anchor="w")

        ctk.CTkLabel(
            left,
            text=f"Phiên bản hiện tại: {APP_VERSION}. Bấm cập nhật để nâng cấp tự động.",
            font=ctk.CTkFont(family="Segoe UI", size=10),
            text_color="#515154"
        ).pack(anchor="w", pady=(2, 0))

        btn_box = ctk.CTkFrame(banner, fg_color="transparent")
        btn_box.pack(side="right", padx=14, pady=10)

        self.btn_update = ctk.CTkButton(
            btn_box,
            text="Cập Nhật Ngay",
            font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
            fg_color="#0071E3",
            hover_color="#0077ED",
            text_color="#FFFFFF",
            corner_radius=8,
            width=120,
            height=30,
            command=lambda: self.start_download_update(exe_url, latest_tag)
        )
        self.btn_update.pack(side="right")

    def start_download_update(self, exe_url, latest_tag):
        self.btn_update.configure(state="disabled", text="Đang tải 0%...")
        threading.Thread(
            target=self._download_and_install_update,
            args=(exe_url, latest_tag),
            daemon=True
        ).start()

    def _download_and_install_update(self, exe_url, latest_tag):
        try:
            temp_dir = tempfile.gettempdir()
            target_exe_name = f"ClaimHelper_update_{int(time.time())}.exe"
            temp_file = os.path.join(temp_dir, target_exe_name)

            req = urllib.request.Request(exe_url, headers={"User-Agent": "ClaimHelper-App"})
            with urllib.request.urlopen(req, timeout=60) as resp:
                total_size = int(resp.headers.get("content-length", 0))
                downloaded = 0
                chunk_size = 64 * 1024
                with open(temp_file, "wb") as f:
                    while True:
                        chunk = resp.read(chunk_size)
                        if not chunk:
                            break
                        f.write(chunk)
                        downloaded += len(chunk)
                        if total_size > 0:
                            pct = int(downloaded / total_size * 100)
                            self.after(0, self.btn_update.configure, {"text": f"Đang tải {pct}%..."})

            self.after(0, self.btn_update.configure, {"text": "Đang khởi động..."})

            # Determine destination path
            is_frozen = getattr(sys, "frozen", False)
            if is_frozen:
                current_exe = os.path.abspath(sys.executable)
            else:
                current_exe = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ClaimHelper.exe")

            updater_bat = os.path.join(temp_dir, f"claim_update_{int(time.time())}.bat")
            bat_script = f"""@echo off
timeout /t 1 /nobreak > nul
move /y "{temp_file}" "{current_exe}" > nul
start "" "{current_exe}"
del "%~f0"
exit
"""
            with open(updater_bat, "w", encoding="utf-8") as bf:
                bf.write(bat_script)

            subprocess.Popen(["cmd.exe", "/c", updater_bat], creationflags=0x08000000)
            self.after(100, self.destroy)
        except Exception as e:
            self.after(0, self.on_update_failed, str(e))

    def on_update_failed(self, err):
        if hasattr(self, 'btn_update'):
            self.btn_update.configure(state="normal", text="Thử lại")
        messagebox.showerror("Cập nhật thất bại", f"Không thể tải bản cập nhật: {err}")

    def setup_ui(self):
        # 1. Header Frame
        header = ctk.CTkFrame(self, fg_color="transparent")
        header.pack(fill="x", padx=32, pady=(25, 6))

        title_row = ctk.CTkFrame(header, fg_color="transparent")
        title_row.pack(anchor="w")

        title_lbl = ctk.CTkLabel(
            title_row,
            text="Claim Helper",
            font=ctk.CTkFont(family="Segoe UI", size=26, weight="bold"),
            text_color="#1D1D1F"
        )
        title_lbl.pack(side="left")

        ctk.CTkLabel(
            title_row,
            text=f"  {APP_VERSION}  ",
            font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold"),
            fg_color="#EBF3FE",
            text_color="#0071E3",
            corner_radius=6
        ).pack(side="left", padx=(10, 0), pady=(6, 0))

        sub_lbl = ctk.CTkLabel(
            header,
            text="Hệ thống tự động đối chiếu & điền Purchase Cost cho SOFACOMPANY",
            font=ctk.CTkFont(family="Segoe UI", size=12),
            text_color="#86868B"
        )
        sub_lbl.pack(anchor="w", pady=(2, 0))

        # Dynamic Update Banner Container (shows up if new version is found on GitHub)
        self.update_banner_container = ctk.CTkFrame(self, fg_color="transparent")
        self.update_banner_container.pack(fill="x", padx=32, pady=(0, 6))

        # 2. Main Scrollable Container
        self.main_scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self.main_scroll.pack(fill="both", expand=True, padx=32, pady=(0, 15))

        # ========================================================
        # CARD 1: Master Overview (Dedicated Drop & Clear Target)
        # ========================================================
        self.card_ov = ctk.CTkFrame(
            self.main_scroll,
            fg_color="#FFFFFF",
            corner_radius=14,
            border_width=1,
            border_color="#E5E5EA"
        )
        self.card_ov.pack(fill="x", pady=(0, 14))

        ov_top = ctk.CTkFrame(self.card_ov, fg_color="transparent")
        ov_top.pack(fill="x", padx=18, pady=(14, 8))

        ctk.CTkLabel(
            ov_top,
            text="📊  FILE MASTER OVERVIEW",
            font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
            text_color="#0071E3"
        ).pack(side="left")

        ov_btn_box = ctk.CTkFrame(ov_top, fg_color="transparent")
        ov_btn_box.pack(side="right")

        self.btn_refresh_cache = ctk.CTkButton(
            ov_btn_box,
            text="Đọc Lại Master",
            font=ctk.CTkFont(family="Segoe UI", size=11),
            fg_color="#F2F2F7",
            text_color="#1D1D1F",
            hover_color="#E5E5EA",
            corner_radius=8,
            height=26,
            command=self.force_reload_master
        )

        self.btn_clear_ov = ctk.CTkButton(
            ov_btn_box,
            text="Xóa File",
            font=ctk.CTkFont(family="Segoe UI", size=11),
            fg_color="transparent",
            text_color="#FF3B30",
            hover_color="#FFECEB",
            corner_radius=8,
            height=26,
            command=self.clear_overview
        )

        self.btn_pick_ov = ctk.CTkButton(
            ov_btn_box,
            text="Chọn File Overview...",
            font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
            fg_color="#F2F2F7",
            text_color="#0071E3",
            hover_color="#E5E5EA",
            corner_radius=8,
            height=26,
            command=self.browse_overview
        )
        self.btn_pick_ov.pack(side="right", padx=(6, 0))

        # Container for Master file display or Drop hint
        self.ov_display_container = ctk.CTkFrame(self.card_ov, fg_color="#F9F9FB", corner_radius=10)
        self.ov_display_container.pack(fill="x", padx=16, pady=(0, 14))

        self.render_overview_display()

        # ========================================================
        # CARD 2: Claim Versus Files (Multi-file select & Drop Target)
        # ========================================================
        self.card_claims = ctk.CTkFrame(
            self.main_scroll,
            fg_color="#FFFFFF",
            corner_radius=14,
            border_width=1,
            border_color="#E5E5EA"
        )
        self.card_claims.pack(fill="x", pady=(0, 14))

        claim_top = ctk.CTkFrame(self.card_claims, fg_color="transparent")
        claim_top.pack(fill="x", padx=18, pady=(14, 8))

        self.claim_count_lbl = ctk.CTkLabel(
            claim_top,
            text=f"📋  DANH SÁCH FILE CLAIM VERSUS ({len(self.claim_files)})",
            font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
            text_color="#0071E3"
        )
        self.claim_count_lbl.pack(side="left")

        actions_box = ctk.CTkFrame(claim_top, fg_color="transparent")
        actions_box.pack(side="right")

        ctk.CTkButton(
            actions_box,
            text="+ Chọn File(s)...",
            font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
            fg_color="#F2F2F7",
            text_color="#0071E3",
            hover_color="#E5E5EA",
            corner_radius=8,
            height=26,
            command=self.browse_claim_files
        ).pack(side="left", padx=4)

        ctk.CTkButton(
            actions_box,
            text="📂 Thư Mục...",
            font=ctk.CTkFont(family="Segoe UI", size=11),
            fg_color="#F2F2F7",
            text_color="#1D1D1F",
            hover_color="#E5E5EA",
            corner_radius=8,
            height=26,
            command=self.browse_claim_folder
        ).pack(side="left", padx=4)

        ctk.CTkButton(
            actions_box,
            text="Xóa Hết",
            font=ctk.CTkFont(family="Segoe UI", size=11),
            fg_color="transparent",
            text_color="#FF3B30",
            hover_color="#FFECEB",
            corner_radius=8,
            height=26,
            command=self.clear_all_claims
        ).pack(side="left", padx=4)

        # File List Inner Container
        self.file_list_frame = ctk.CTkFrame(self.card_claims, fg_color="#F9F9FB", corner_radius=10)
        self.file_list_frame.pack(fill="x", padx=16, pady=(0, 14))

        self.render_file_list()

        # ========================================================
        # ACTION: Start Processing Button & Progress Bar
        # ========================================================
        act_box = ctk.CTkFrame(self.main_scroll, fg_color="transparent")
        act_box.pack(fill="x", pady=(4, 12))

        self.btn_run = ctk.CTkButton(
            act_box,
            text="Bắt Đầu Đối Chiếu & Điền Giá (Tạo File _filled)",
            font=ctk.CTkFont(family="Segoe UI", size=14, weight="bold"),
            fg_color="#0071E3",
            hover_color="#0077ED",
            text_color="#FFFFFF",
            corner_radius=22,
            height=46,
            command=self.start_processing
        )
        self.btn_run.pack(fill="x")

        self.prog_bar = ctk.CTkProgressBar(
            self.main_scroll,
            progress_color="#0071E3",
            fg_color="#E5E5EA",
            height=4,
            corner_radius=2
        )
        self.prog_bar.set(0)
        self.prog_bar.pack(fill="x", pady=(0, 12))

        # ========================================================
        # CARD 3: Apple-Style Live Activity & Results Card
        # ========================================================
        self.card_results = ctk.CTkFrame(
            self.main_scroll,
            fg_color="#FFFFFF",
            corner_radius=14,
            border_width=1,
            border_color="#E5E5EA"
        )
        self.card_results.pack(fill="both", expand=True, pady=(0, 6))

        results_head = ctk.CTkFrame(self.card_results, fg_color="transparent")
        results_head.pack(fill="x", padx=18, pady=(14, 8))

        ctk.CTkLabel(
            results_head,
            text="⚡  TIẾN ĐỘ & KẾT QUẢ",
            font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
            text_color="#0071E3"
        ).pack(side="left")

        self.status_badge = ctk.CTkLabel(
            results_head,
            text="Sẵn sàng",
            font=ctk.CTkFont(family="Segoe UI", size=11),
            text_color="#86868B"
        )
        self.status_badge.pack(side="right")

        # Container for concise result items
        self.results_feed = ctk.CTkFrame(self.card_results, fg_color="#F9F9FB", corner_radius=10)
        self.results_feed.pack(fill="both", expand=True, padx=16, pady=(0, 14))

        self.render_initial_feed()

    # ----------------------------------------------------
    # UI Renderers
    # ----------------------------------------------------
    def render_initial_feed(self):
        for child in self.results_feed.winfo_children():
            child.destroy()
        
        hint = ctk.CTkLabel(
            self.results_feed,
            text="Kết quả xử lý sẽ hiển thị tại đây một cách ngắn gọn, rõ ràng.",
            font=ctk.CTkFont(family="Segoe UI", size=12),
            text_color="#86868B",
            pady=20
        )
        hint.pack(fill="both", expand=True)

    def render_overview_display(self):
        for child in self.ov_display_container.winfo_children():
            child.destroy()

        if self.overview_path and os.path.exists(self.overview_path):
            self.btn_clear_ov.pack(side="right", padx=(4, 0))
            self.btn_refresh_cache.pack(side="right", padx=(4, 0))
            size_str = format_file_size(os.path.getsize(self.overview_path))
            
            # Check cache status
            c_path = get_cache_path()
            has_cache = False
            if os.path.exists(c_path):
                try:
                    with open(c_path, 'rb') as f:
                        meta = pickle.load(f)
                    if (meta.get('file_path') == os.path.abspath(self.overview_path) and
                        meta.get('file_size') == os.path.getsize(self.overview_path)):
                        has_cache = True
                except Exception:
                    pass

            row = ctk.CTkFrame(self.ov_display_container, fg_color="#FFFFFF", corner_radius=8)
            row.pack(fill="x", padx=8, pady=8)

            ctk.CTkLabel(
                row,
                text="📊",
                font=ctk.CTkFont(size=14)
            ).pack(side="left", padx=(10, 8), pady=8)

            info_box = ctk.CTkFrame(row, fg_color="transparent")
            info_box.pack(side="left", fill="x", expand=True, pady=6)

            title_row = ctk.CTkFrame(info_box, fg_color="transparent")
            title_row.pack(anchor="w")

            ctk.CTkLabel(
                title_row,
                text=os.path.basename(self.overview_path),
                font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
                text_color="#1D1D1F"
            ).pack(side="left")

            if has_cache:
                ctk.CTkLabel(
                    title_row,
                    text="  ✓ ĐÃ GHI NHỚ  ",
                    font=ctk.CTkFont(family="Segoe UI", size=9, weight="bold"),
                    fg_color="#EBF8F2",
                    text_color="#16A34A",
                    corner_radius=4
                ).pack(side="left", padx=(8, 0))

            ctk.CTkLabel(
                info_box,
                text=f"{self.overview_path}  •  {size_str}",
                font=ctk.CTkFont(family="Segoe UI", size=10),
                text_color="#86868B",
                anchor="w"
            ).pack(anchor="w", pady=(2, 0))

            # Inline Delete Button
            ctk.CTkButton(
                row,
                text="✕",
                font=ctk.CTkFont(size=11, weight="bold"),
                fg_color="transparent",
                text_color="#86868B",
                hover_color="#FFECEB",
                corner_radius=6,
                width=24,
                height=24,
                command=self.clear_overview
            ).pack(side="right", padx=(4, 10), pady=8)
        else:
            self.btn_clear_ov.pack_forget()
            self.btn_refresh_cache.pack_forget()
            drop_hint = ctk.CTkLabel(
                self.ov_display_container,
                text="📥  Kéo thả file Master Overview (.xlsx) vào đây\nhoặc bấm nút 'Chọn File Overview...' ở góc trên",
                font=ctk.CTkFont(family="Segoe UI", size=12),
                text_color="#86868B",
                pady=18
            )
            drop_hint.pack(fill="both", expand=True)

    def clear_overview(self):
        self.overview_path = ""
        self.save_state()
        self.render_overview_display()

    def render_file_list(self):
        for child in self.file_list_frame.winfo_children():
            child.destroy()

        if not self.claim_files:
            empty_lbl = ctk.CTkLabel(
                self.file_list_frame,
                text="📥  Kéo thả một hoặc nhiều file Claim Versus (.xlsx) vào đây\nhoặc bấm nút '+ Chọn File(s)...' ở góc trên",
                font=ctk.CTkFont(family="Segoe UI", size=12),
                text_color="#86868B",
                pady=22
            )
            empty_lbl.pack(fill="both", expand=True)
            self.claim_count_lbl.configure(text="📋  DANH SÁCH FILE CLAIM VERSUS (0)")
            return

        self.claim_count_lbl.configure(text=f"📋  DANH SÁCH FILE CLAIM VERSUS ({len(self.claim_files)})")

        for idx, fpath in enumerate(self.claim_files):
            fname = os.path.basename(fpath)
            is_casa = "casa" in fname.lower()
            tag_text = "Casa (CS)" if is_casa else "Nhan Hoang (NH)"
            tag_color = "#EBF3FE" if is_casa else "#EBF8F2"
            tag_text_color = "#0071E3" if is_casa else "#16A34A"

            base, ext = os.path.splitext(fname)
            out_preview = f"{base}_filled{ext}" if not base.endswith("_filled") else fname

            row = ctk.CTkFrame(self.file_list_frame, fg_color="#FFFFFF", corner_radius=8)
            row.pack(fill="x", padx=8, pady=4)

            # Icon & Name
            ctk.CTkLabel(
                row,
                text="📄",
                font=ctk.CTkFont(size=13)
            ).pack(side="left", padx=(10, 6), pady=6)

            info = ctk.CTkFrame(row, fg_color="transparent")
            info.pack(side="left", fill="x", expand=True, pady=4)

            ctk.CTkLabel(
                info,
                text=fname,
                font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
                text_color="#1D1D1F",
                anchor="w"
            ).pack(anchor="w")

            ctk.CTkLabel(
                info,
                text=f"↳ File xuất sẽ tạo: {out_preview}",
                font=ctk.CTkFont(family="Segoe UI", size=10),
                text_color="#86868B",
                anchor="w"
            ).pack(anchor="w")

            # Badge
            badge = ctk.CTkLabel(
                row,
                text=f"  {tag_text}  ",
                font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold"),
                fg_color=tag_color,
                text_color=tag_text_color,
                corner_radius=6
            )
            badge.pack(side="left", padx=8, pady=6)

            # Delete button
            del_btn = ctk.CTkButton(
                row,
                text="✕",
                font=ctk.CTkFont(size=11, weight="bold"),
                fg_color="transparent",
                text_color="#86868B",
                hover_color="#FFECEB",
                corner_radius=6,
                width=24,
                height=24,
                command=lambda p=fpath: self.remove_claim_file(p)
            )
            del_btn.pack(side="right", padx=(4, 8), pady=6)

    def remove_claim_file(self, fpath):
        if fpath in self.claim_files:
            self.claim_files.remove(fpath)
            self.save_state()
            self.render_file_list()

    def clear_all_claims(self):
        if not self.claim_files:
            return
        if messagebox.askyesno("Xác nhận", "Bạn có chắc muốn xóa tất cả file khỏi danh sách không?"):
            self.claim_files.clear()
            self.save_state()
            self.render_file_list()

    # ----------------------------------------------------
    # Drag & Drop Handlers
    # ----------------------------------------------------
    def setup_drag_and_drop(self):
        targets = [self.card_ov, self.ov_display_container, self.card_claims, self.file_list_frame, self]
        for t in targets:
            try:
                t.drop_target_register(tkdnd.DND_FILES)
                t.dnd_bind('<<Drop>>', self.on_drop_generic)
            except Exception:
                pass

        try:
            self.card_ov.dnd_bind('<<Drop>>', self.on_drop_overview)
            self.ov_display_container.dnd_bind('<<Drop>>', self.on_drop_overview)
            self.card_ov.dnd_bind('<<DropEnter>>', lambda e: self.card_ov.configure(border_color="#0071E3", fg_color="#F0F8FF"))
            self.card_ov.dnd_bind('<<DropLeave>>', lambda e: self.card_ov.configure(border_color="#E5E5EA", fg_color="#FFFFFF"))
        except Exception:
            pass

        try:
            self.card_claims.dnd_bind('<<Drop>>', self.on_drop_claims)
            self.file_list_frame.dnd_bind('<<Drop>>', self.on_drop_claims)
            self.card_claims.dnd_bind('<<DropEnter>>', lambda e: self.card_claims.configure(border_color="#0071E3", fg_color="#F0F8FF"))
            self.card_claims.dnd_bind('<<DropLeave>>', lambda e: self.card_claims.configure(border_color="#E5E5EA", fg_color="#FFFFFF"))
        except Exception:
            pass

    def on_drop_overview(self, event):
        self.card_ov.configure(border_color="#E5E5EA", fg_color="#FFFFFF")
        paths = parse_drop_paths(event.data)
        for p in paths:
            if os.path.isfile(p) and p.endswith(".xlsx"):
                self.overview_path = p
                self.save_state()
                self.render_overview_display()
                return

    def on_drop_claims(self, event):
        self.card_claims.configure(border_color="#E5E5EA", fg_color="#FFFFFF")
        paths = parse_drop_paths(event.data)
        self.add_claim_paths(paths)

    def on_drop_generic(self, event):
        self.card_ov.configure(border_color="#E5E5EA", fg_color="#FFFFFF")
        self.card_claims.configure(border_color="#E5E5EA", fg_color="#FFFFFF")
        paths = parse_drop_paths(event.data)
        if not paths:
            return

        claims_to_add = []
        for p in paths:
            if os.path.isfile(p) and p.endswith(".xlsx"):
                fname = os.path.basename(p).lower()
                if "overview" in fname:
                    self.overview_path = p
                    self.save_state()
                    self.render_overview_display()
                else:
                    claims_to_add.append(p)
            elif os.path.isdir(p):
                claims_to_add.append(p)

        if claims_to_add:
            self.add_claim_paths(claims_to_add)

    def add_claim_paths(self, paths):
        added = 0
        for p in paths:
            if os.path.isfile(p) and p.endswith(".xlsx") and not p.endswith("_backup.xlsx") and not p.startswith("~$"):
                if p not in self.claim_files:
                    self.claim_files.append(p)
                    added += 1
            elif os.path.isdir(p):
                for f in os.listdir(p):
                    if "versus" in f.lower() and f.endswith(".xlsx") and not f.endswith("_backup.xlsx") and not f.endswith("_filled.xlsx") and not f.startswith("~$"):
                        fp = os.path.join(p, f)
                        if fp not in self.claim_files:
                            self.claim_files.append(fp)
                            added += 1
        if added > 0:
            self.save_state()
            self.render_file_list()

    # ----------------------------------------------------
    # File Dialogs
    # ----------------------------------------------------
    def browse_overview(self):
        f = filedialog.askopenfilename(
            title="Chọn file Master Overview",
            filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")]
        )
        if f:
            self.overview_path = os.path.abspath(f)
            self.save_state()
            self.render_overview_display()

    def browse_claim_files(self):
        files = filedialog.askopenfilenames(
            title="Chọn một hoặc nhiều file Claim Versus",
            filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")]
        )
        if files:
            self.add_claim_paths([os.path.abspath(f) for f in files])

    def browse_claim_folder(self):
        d = filedialog.askdirectory(title="Chọn thư mục chứa các file Claim Versus")
        if d:
            self.add_claim_paths([d])

    def force_reload_master(self):
        if not self.overview_path or not os.path.exists(self.overview_path):
            messagebox.showwarning("Cảnh báo", "Vui lòng chọn file Master Overview trước!")
            return
        
        c_path = get_cache_path()
        if os.path.exists(c_path):
            try:
                os.remove(c_path)
            except Exception:
                pass
        self.render_overview_display()
        messagebox.showinfo("Thông báo", "Đã xóa bộ nhớ ghi nhớ! Lần chạy tới sẽ đọc lại file Master từ đầu.")

    def open_current_folder(self):
        if self.last_created_files:
            reveal_in_explorer(self.last_created_files[0])
        elif self.claim_files:
            folder = os.path.dirname(self.claim_files[0])
            os.startfile(folder)
        elif self.overview_path:
            folder = os.path.dirname(self.overview_path)
            os.startfile(folder)
        else:
            os.startfile(os.path.abspath(os.path.dirname(__file__)))

    # ----------------------------------------------------
    # Thread-Safe Apple Results Feed Updates
    # ----------------------------------------------------
    def append_feed_item(self, filename, matched, total, rate_str, full_fpath=None):
        row = ctk.CTkFrame(self.results_feed, fg_color="#FFFFFF", corner_radius=8)
        row.pack(fill="x", padx=8, pady=3)

        ctk.CTkLabel(
            row,
            text="✓",
            font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
            text_color="#16A34A"
        ).pack(side="left", padx=(10, 8), pady=6)

        ctk.CTkLabel(
            row,
            text=filename,
            font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
            text_color="#1D1D1F"
        ).pack(side="left", pady=6)

        ctk.CTkLabel(
            row,
            text=f"•  Khớp {matched}/{total} dòng ({rate_str})",
            font=ctk.CTkFont(family="Segoe UI", size=11),
            text_color="#86868B"
        ).pack(side="left", padx=8, pady=6)

        if full_fpath and os.path.exists(full_fpath):
            ctk.CTkButton(
                row,
                text="📂 Xem",
                font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold"),
                fg_color="#F2F2F7",
                hover_color="#E5E5EA",
                text_color="#0071E3",
                corner_radius=6,
                width=52,
                height=22,
                command=lambda p=full_fpath: reveal_in_explorer(p)
            ).pack(side="right", padx=(4, 8), pady=6)

    def set_status_text(self, text, color="#86868B"):
        self.status_badge.configure(text=text, text_color=color)

    # ----------------------------------------------------
    # Core Engine Processing (Thread-Safe & Lag-Free)
    # ----------------------------------------------------
    def start_processing(self):
        if not self.overview_path or not os.path.isfile(self.overview_path):
            messagebox.showerror("Thiếu thông tin", "Vui lòng chọn hoặc kéo thả file Master Overview (.xlsx)!")
            return
        if not self.claim_files:
            messagebox.showerror("Thiếu thông tin", "Danh sách file Claim đang trống. Vui lòng chọn hoặc kéo thả ít nhất 1 file Claim Versus!")
            return

        self.btn_run.configure(state="disabled", text="Đang xử lý dữ liệu...")
        self.prog_bar.set(0.05)
        self.set_status_text("● Đang xử lý...", "#0071E3")
        self.last_created_files = []

        # Clear feed
        for child in self.results_feed.winfo_children():
            child.destroy()

        # Run background thread
        thread = threading.Thread(target=self.run_engine, daemon=True)
        thread.start()

    def run_engine(self):
        try:
            # 1. Load Master Overview with Cache
            c_path = get_cache_path()
            ov_abs = os.path.abspath(self.overview_path)
            cur_size = os.path.getsize(ov_abs)
            cur_mtime = os.path.getmtime(ov_abs)

            loaded_from_cache = False
            cs_records = []
            nh_records = []

            if os.path.exists(c_path):
                try:
                    with open(c_path, 'rb') as f:
                        cache_obj = pickle.load(f)
                    if (cache_obj.get('file_path') == ov_abs and
                        cache_obj.get('file_size') == cur_size and
                        abs(cache_obj.get('file_mtime', 0) - cur_mtime) < 2):
                        cs_records = cache_obj['cs_records']
                        nh_records = cache_obj['nh_records']
                        loaded_from_cache = True
                except Exception:
                    pass

            if not loaded_from_cache:
                self.after(0, self.set_status_text, "● Đang đọc file Master...", "#0071E3")
                wb_o = openpyxl.load_workbook(self.overview_path, read_only=True, data_only=True)
                s_o = wb_o['Overview'] if 'Overview' in wb_o.sheetnames else wb_o.active
                headers_o = next(s_o.iter_rows(max_row=1, values_only=True))
                col_idx = {h: i for i, h in enumerate(headers_o)}

                for r in s_o.iter_rows(min_row=2, values_only=True):
                    man = str(r[col_idx.get('Manufacturer', 0)] or '').strip().upper()
                    sup = str(r[col_idx.get('Supplier', 0)] or '').strip().upper()
                    cost = r[col_idx.get('Purchase cost', 0)]
                    note_p = str(r[col_idx.get('Item PO note', 0)] or '').lower()
                    note_s = str(r[col_idx.get('Note (sales)', 0)] or '').lower()

                    if 'box 1 of 2' in note_p or 'box 2 of 2' in note_p or 'box 1 of' in note_s:
                        continue
                    if cost is None or float(cost) <= 0:
                        continue

                    rec = {
                        'art': str(r[col_idx.get('Customer art No', 0)] or '').strip(),
                        'item_name': str(r[col_idx.get('Customer item name', 0)] or '').strip(),
                        'fabric': str(r[col_idx.get('Fabric type', 0)] or '').strip(),
                        'cost': float(cost),
                        'po_date': r[col_idx.get('P.O issued date', 0)],
                        'model': str(r[col_idx.get('Model', 0)] or '').strip()
                    }

                    if man in ['CS', 'WORKSHOP'] or 'CASA' in man or 'CASA' in sup:
                        cs_records.append(rec)
                    elif man in ['NH', 'NHF', 'NH FSC', 'NHAN HOANG'] or 'NHAN HOANG' in sup or 'NH ' in sup:
                        nh_records.append(rec)

                wb_o.close()

                # Save cache
                try:
                    cache_obj = {
                        'file_path': ov_abs,
                        'file_size': cur_size,
                        'file_mtime': cur_mtime,
                        'cached_at': datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                        'cs_records': cs_records,
                        'nh_records': nh_records
                    }
                    with open(c_path, 'wb') as f:
                        pickle.dump(cache_obj, f, protocol=pickle.HIGHEST_PROTOCOL)
                except Exception:
                    pass

            self.after(0, self.render_overview_display)

            # 2. Process Files
            total_files = len(self.claim_files)
            created_files = []

            for i, fpath in enumerate(self.claim_files, start=1):
                fname = os.path.basename(fpath)
                is_casa = "casa" in fname.lower()
                ref_db = cs_records if is_casa else nh_records

                wb = openpyxl.load_workbook(fpath)
                s = wb['Export'] if 'Export' in wb.sheetnames else wb.active
                headers = [cell.value for cell in s[1]]

                # Add Purchase Cost (USD)
                if 'Purchase Cost (USD)' in headers:
                    col_cost_idx = headers.index('Purchase Cost (USD)') + 1
                else:
                    col_cost_idx = len(headers) + 1
                    headers.append('Purchase Cost (USD)')
                    c_h1 = s.cell(row=1, column=col_cost_idx, value='Purchase Cost (USD)')
                    sample_h = s.cell(row=1, column=headers.index('Unit Cost x QTY')+1 if 'Unit Cost x QTY' in headers else 1)
                    if sample_h.has_style:
                        c_h1.font, c_h1.border, c_h1.fill, c_h1.alignment = copy(sample_h.font), copy(sample_h.border), copy(sample_h.fill), copy(sample_h.alignment)

                # Add Purchase Cost Note
                if 'Purchase Cost Note' in headers:
                    col_note_idx = headers.index('Purchase Cost Note') + 1
                else:
                    col_note_idx = len(headers) + 1
                    headers.append('Purchase Cost Note')
                    c_h2 = s.cell(row=1, column=col_note_idx, value='Purchase Cost Note')
                    sample_h = s.cell(row=1, column=headers.index('Unit Cost x QTY')+1 if 'Unit Cost x QTY' in headers else 1)
                    if sample_h.has_style:
                        c_h2.font, c_h2.border, c_h2.fill, c_h2.alignment = copy(sample_h.font), copy(sample_h.border), copy(sample_h.fill), copy(sample_h.alignment)

                cost_letter = get_column_letter(col_cost_idx)
                note_letter = get_column_letter(col_note_idx)
                s.column_dimensions[cost_letter].width = 22
                s.column_dimensions[note_letter].width = 85

                max_row = s.max_row
                art_col_idx = headers.index('No') + 1 if 'No' in headers else 7
                unit_cost_col_idx = headers.index('Unit Cost x QTY') + 1 if 'Unit Cost x QTY' in headers else None

                last_art = s.cell(row=max_row, column=art_col_idx).value
                last_is_total = (last_art is None or str(last_art).strip() == '')
                data_end = max_row - 1 if last_is_total else max_row
                data_count = data_end - 1

                rate_samples = []
                if unit_cost_col_idx:
                    for r_idx in range(2, data_end + 1):
                        art = str(s.cell(row=r_idx, column=art_col_idx).value or '').strip()
                        dkk = s.cell(row=r_idx, column=unit_cost_col_idx).value
                        exact = [rec for rec in ref_db if rec['art'] == art]
                        if exact and isinstance(dkk, (int, float)) and dkk > 0:
                            avg_usd = sum(e['cost'] for e in exact) / len(exact)
                            rate_samples.append(dkk / avg_usd)

                median_rate = sorted(rate_samples)[len(rate_samples)//2] if rate_samples else (8.0087 if is_casa else 7.8210)

                file_matched = 0
                for r_idx in range(2, data_end + 1):
                    art = str(s.cell(row=r_idx, column=art_col_idx).value or '').strip()
                    desc = str(s.cell(row=r_idx, column=headers.index('Description')+1).value or '').strip() if 'Description' in headers else ''
                    desc2 = str(s.cell(row=r_idx, column=headers.index('Description_2')+1).value or '').strip() if 'Description_2' in headers else ''
                    posting_date = s.cell(row=r_idx, column=headers.index('PostingDate')+1).value if 'PostingDate' in headers else None
                    days = s.cell(row=r_idx, column=headers.index('Days Between Order And Complaint')+1).value if 'Days Between Order And Complaint' in headers else 0
                    days = days if isinstance(days, (int, float)) else 0
                    unit_cost_dkk = s.cell(row=r_idx, column=unit_cost_col_idx).value if unit_cost_col_idx else 0
                    unit_cost_dkk = unit_cost_dkk if isinstance(unit_cost_dkk, (int, float)) else 0

                    order_date = None
                    if isinstance(posting_date, datetime.datetime):
                        order_date = posting_date - datetime.timedelta(days=int(days))
                    elif isinstance(posting_date, datetime.date):
                        order_date = posting_date - datetime.timedelta(days=int(days))

                    exact_matches = [rec for rec in ref_db if rec['art'] == art]
                    selected_cost = None
                    note = ""

                    if exact_matches:
                        file_matched += 1
                        if unit_cost_dkk > 0:
                            target_usd = unit_cost_dkk / median_rate
                            plausible = [e for e in exact_matches if 0.7 * target_usd <= e['cost'] <= 1.3 * target_usd]
                            if not plausible: plausible = exact_matches
                        else:
                            plausible = exact_matches

                        valid_dates = [e for e in plausible if isinstance(e['po_date'], (datetime.datetime, datetime.date))]
                        if valid_dates and order_date:
                            o_d = order_date.date() if isinstance(order_date, datetime.datetime) else order_date
                            best_match = min(valid_dates, key=lambda e: abs((e['po_date'].date() if isinstance(e['po_date'], datetime.datetime) else e['po_date']) - o_d))
                            selected_cost = best_match['cost']
                            po_str = best_match['po_date'].strftime('%d/%m/%Y')
                            od_str = order_date.strftime('%d/%m/%Y')
                            note = f"Khớp 100% Art No ({art}). Chọn {selected_cost}$ theo đợt PO ({po_str}) gần nhất ngày đặt hàng ({od_str})."
                        else:
                            best_match = min(plausible, key=lambda e: abs(e['cost'] - (unit_cost_dkk/median_rate if unit_cost_dkk > 0 else plausible[0]['cost'])))
                            selected_cost = best_match['cost']
                            note = f"Khớp 100% Art No ({art}). Chọn {selected_cost}$ từ danh mục trong Overview."
                    else:
                        candidates = [rec for rec in ref_db if rec['art'].startswith(art) or art in rec['art']]
                        if not candidates and len(art) >= 6:
                            candidates = [rec for rec in ref_db if rec['art'].startswith(art[:6])]

                        if candidates and unit_cost_dkk > 0:
                            file_matched += 1
                            est_usd = unit_cost_dkk / median_rate
                            best_cand = min(candidates, key=lambda c: abs(c['cost'] - est_usd))
                            selected_cost = best_cand['cost']
                            note = f"Mã rút gọn MM ({art}). Dựa vào Unit Cost {unit_cost_dkk:.2f} DKK (~{est_usd:.2f}$), khớp phân khúc vải '{best_cand['fabric']}' (giá {selected_cost}$)."
                        elif candidates:
                            file_matched += 1
                            selected_cost = round(sum(c['cost'] for c in candidates) / len(candidates), 2)
                            note = f"Mã rút gọn MM ({art}). Lấy giá trung bình phân khúc {selected_cost}$."
                        elif unit_cost_dkk > 0:
                            file_matched += 1
                            selected_cost = round(unit_cost_dkk / median_rate, 2)
                            note = f"Mã {art} không còn trong Overview. Tính theo tỷ giá quy đổi Unit Cost: {unit_cost_dkk:.2f} DKK / {median_rate:.4f} = {selected_cost}$."
                        else:
                            selected_cost = 0.0
                            note = f"Không tìm thấy mã {art} trong Overview."

                    c_cost = s.cell(row=r_idx, column=col_cost_idx, value=selected_cost)
                    c_cost.number_format = '#,##0.00'
                    sample_c = s.cell(row=r_idx, column=unit_cost_col_idx if unit_cost_col_idx else 1)
                    if sample_c.has_style:
                        c_cost.font, c_cost.border = copy(sample_c.font), copy(sample_c.border)

                    c_note = s.cell(row=r_idx, column=col_note_idx, value=note)
                    if sample_c.has_style:
                        c_note.font, c_note.border = copy(sample_c.font), copy(sample_c.border)

                # Total Row
                if last_is_total:
                    c_total = s.cell(row=max_row, column=col_cost_idx, value=f"=SUM({cost_letter}2:{cost_letter}{data_end})")
                    c_total.number_format = '#,##0.00'
                    sample_t = s.cell(row=max_row, column=unit_cost_col_idx if unit_cost_col_idx else 1)
                    if sample_t.has_style:
                        c_total.font, c_total.border = copy(sample_t.font), copy(sample_t.border)

                    c_t_note = s.cell(row=max_row, column=col_note_idx, value=f"Tổng Purchase Cost (USD) của {data_count} vụ khiếu nại")
                    if sample_t.has_style:
                        c_t_note.font, c_t_note.border = copy(sample_t.font), copy(sample_t.border)

                # Save as adjacent file with '_filled' suffix
                dir_name = os.path.dirname(fpath)
                base_name, ext = os.path.splitext(fname)
                if base_name.endswith("_filled"):
                    out_fpath = fpath
                else:
                    out_fpath = os.path.join(dir_name, f"{base_name}_filled{ext}")

                wb.save(out_fpath)
                wb.close()

                out_abs = os.path.abspath(out_fpath)
                created_files.append(out_abs)

                rate_str = f"{(file_matched / data_count * 100):.1f}%" if data_count > 0 else "100%"
                out_name = os.path.basename(out_fpath)
                self.after(0, self.append_feed_item, out_name, file_matched, data_count, rate_str, out_abs)

                progress_val = 0.1 + 0.9 * (i / total_files)
                self.after(0, self.prog_bar.set, progress_val)

            # Completion
            self.last_created_files = created_files
            self.after(0, self.finish_processing_apple, True, total_files, created_files)

        except Exception as e:
            self.after(0, self.finish_processing_apple, False, 0, [], str(e))

    def finish_processing_apple(self, success, total_files, created_files, err_msg=""):
        self.prog_bar.set(1.0 if success else 0)
        self.btn_run.configure(state="normal", text="Bắt Đầu Đối Chiếu & Điền Giá (Tạo File _filled)")

        if success:
            self.set_status_text("✓ Hoàn tất", "#16A34A")

            # Determine subtitle message with full path
            if len(created_files) == 1:
                sub_text = f"Đã tạo file: {created_files[0]}"
                target_file_to_reveal = created_files[0]
            else:
                lines = [f"• {p}" for p in created_files]
                sub_text = f"Đã tạo {len(created_files)} file:\n" + "\n".join(lines)
                target_file_to_reveal = created_files[0] if created_files else ""

            # Render completion banner smoothly without blocking modal
            banner = ctk.CTkFrame(self.results_feed, fg_color="#EBF8F2", corner_radius=10)
            banner.pack(fill="x", padx=8, pady=(8, 4))

            b_left = ctk.CTkFrame(banner, fg_color="transparent")
            b_left.pack(side="left", fill="x", expand=True, padx=14, pady=12)

            ctk.CTkLabel(
                b_left,
                text=f"🎉 Hoàn tất xuất sắc {total_files} file kết quả!",
                font=ctk.CTkFont(family="Segoe UI", size=13, weight="bold"),
                text_color="#16A34A"
            ).pack(anchor="w")

            ctk.CTkLabel(
                b_left,
                text=sub_text,
                font=ctk.CTkFont(family="Segoe UI", size=11),
                text_color="#1F2937",
                wraplength=520,
                justify="left"
            ).pack(anchor="w", pady=(3, 0))

            ctk.CTkButton(
                banner,
                text="📂 Xem File",
                font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
                fg_color="#16A34A",
                text_color="#FFFFFF",
                hover_color="#15803D",
                corner_radius=8,
                height=32,
                command=lambda p=target_file_to_reveal: reveal_in_explorer(p)
            ).pack(side="right", padx=14, pady=12)

        else:
            self.set_status_text("❌ Lỗi", "#FF3B30")
            err_box = ctk.CTkFrame(self.results_feed, fg_color="#FFECEB", corner_radius=8)
            err_box.pack(fill="x", padx=8, pady=6)
            ctk.CTkLabel(
                err_box,
                text=f"Đã xảy ra lỗi: {err_msg}",
                font=ctk.CTkFont(family="Segoe UI", size=11),
                text_color="#FF3B30"
            ).pack(padx=10, pady=8)

if __name__ == "__main__":
    app = ClaimHelperAppleApp()
    app.mainloop()
